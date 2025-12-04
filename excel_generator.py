"""
Excel Report Generator for Student Reports
Exports class reports to Excel with proper formatting
"""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import io
from datetime import datetime

def generate_class_report_excel(class_data, term_data, students_data):
    """
    Generate Excel report for entire class
    
    Args:
        class_data: dict with class information
        term_data: dict with term information  
        students_data: list of dicts with student scores
    
    Returns:
        BytesIO object containing Excel file
    """
    wb = Workbook()
    ws = wb.active
    ws.title = f"{class_data['name']} Report"
    
    # Styling
    header_font = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    sub_header_font = Font(name='Calibri', size=11, bold=True)
    sub_header_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    center_align = Alignment(horizontal='center', vertical='center')
    
    # Title Section
    ws.merge_cells('A1:L1')
    title_cell = ws['A1']
    title_cell.value = "SOW THE SEED MODEL COLLEGE"
    title_cell.font = Font(name='Calibri', size=16, bold=True)
    title_cell.alignment = center_align
    
    ws.merge_cells('A2:L2')
    subtitle_cell = ws['A2']
    subtitle_cell.value = f"END OF {term_data['term'].upper()} REPORT - {term_data['academic_year']}"
    subtitle_cell.font = Font(name='Calibri', size=12, bold=True)
    subtitle_cell.alignment = center_align
    
    ws.merge_cells('A3:L3')
    class_cell = ws['A3']
    class_cell.value = f"CLASS: {class_data['name']}"
    class_cell.font = Font(name='Calibri', size=11, bold=True)
    class_cell.alignment = center_align
    
    # Empty row
    current_row = 5
    
    # Get all subjects from first student
    if not students_data or not students_data[0].get('scores'):
        return None
    
    subjects = [score['subjects']['name'] for score in students_data[0]['scores']]
    num_subjects = len(subjects)
    
    # Headers
    headers = ['S/N', 'ADMISSION NO', 'STUDENT NAME']
    
    # Add subject headers
    for subject in subjects:
        headers.extend([
            f"{subject} - CA1",
            f"{subject} - CA2", 
            f"{subject} - EXAM",
            f"{subject} - TOTAL",
            f"{subject} - GRADE"
        ])
    
    headers.extend(['GRAND TOTAL', 'AVERAGE', 'POSITION'])
    
    # Write headers
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=current_row, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = border
    
    current_row += 1
    
    # Write student data
    for student_num, student in enumerate(students_data, 1):
        col = 1
        
        # S/N
        ws.cell(row=current_row, column=col, value=student_num).border = border
        col += 1
        
        # Admission number
        ws.cell(row=current_row, column=col, value=student['student']['admission_number']).border = border
        col += 1
        
        # Student name
        full_name = f"{student['student']['first_name']} {student['student']['last_name']}"
        ws.cell(row=current_row, column=col, value=full_name).border = border
        col += 1
        
        # Scores for each subject
        grand_total = 0
        for score in student['scores']:
            # CA1
            cell = ws.cell(row=current_row, column=col, value=float(score['ca1']))
            cell.border = border
            cell.alignment = center_align
            col += 1
            
            # CA2
            cell = ws.cell(row=current_row, column=col, value=float(score['ca2']))
            cell.border = border
            cell.alignment = center_align
            col += 1
            
            # Exam
            cell = ws.cell(row=current_row, column=col, value=float(score['exam']))
            cell.border = border
            cell.alignment = center_align
            col += 1
            
            # Total
            cell = ws.cell(row=current_row, column=col, value=float(score['total']))
            cell.border = border
            cell.alignment = center_align
            cell.font = Font(bold=True)
            grand_total += float(score['total'])
            col += 1
            
            # Grade
            cell = ws.cell(row=current_row, column=col, value=score['grade'])
            cell.border = border
            cell.alignment = center_align
            cell.font = Font(bold=True)
            
            # Color code grades
            if score['grade'] in ['A+', 'B+']:
                cell.fill = PatternFill(start_color='C6E0B4', end_color='C6E0B4', fill_type='solid')
            elif score['grade'] in ['C', 'D']:
                cell.fill = PatternFill(start_color='FFE699', end_color='FFE699', fill_type='solid')
            elif score['grade'] in ['E', 'F']:
                cell.fill = PatternFill(start_color='F4B084', end_color='F4B084', fill_type='solid')
            
            col += 1
        
        # Grand Total
        cell = ws.cell(row=current_row, column=col, value=grand_total)
        cell.border = border
        cell.alignment = center_align
        cell.font = Font(bold=True, size=11)
        cell.fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
        col += 1
        
        # Average
        average = grand_total / len(student['scores']) if student['scores'] else 0
        cell = ws.cell(row=current_row, column=col, value=round(average, 2))
        cell.border = border
        cell.alignment = center_align
        cell.font = Font(bold=True)
        col += 1
        
        # Position
        position = student.get('statistics', {}).get('position', '-')
        cell = ws.cell(row=current_row, column=col, value=f"{position}{get_ordinal_suffix(position)}")
        cell.border = border
        cell.alignment = center_align
        cell.font = Font(bold=True)
        
        current_row += 1
    
    # Auto-adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    # Freeze panes (freeze headers)
    ws.freeze_panes = 'D6'
    
    # Save to BytesIO
    excel_file = io.BytesIO()
    wb.save(excel_file)
    excel_file.seek(0)
    
    return excel_file


def generate_student_report_excel(student_data):
    """
    Generate Excel report for individual student
    
    Args:
        student_data: dict with student information and scores
    
    Returns:
        BytesIO object containing Excel file
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Student Report"
    
    # Styling
    title_font = Font(name='Calibri', size=14, bold=True)
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    center_align = Alignment(horizontal='center', vertical='center')
    
    # Title
    ws.merge_cells('A1:G1')
    ws['A1'] = "SOW THE SEED MODEL COLLEGE"
    ws['A1'].font = Font(name='Calibri', size=16, bold=True)
    ws['A1'].alignment = center_align
    
    # Student Info
    row = 3
    ws[f'A{row}'] = "Student Name:"
    ws[f'A{row}'].font = Font(bold=True)
    ws[f'B{row}'] = f"{student_data['student']['first_name']} {student_data['student']['last_name']}"
    ws.merge_cells(f'B{row}:D{row}')
    
    row += 1
    ws[f'A{row}'] = "Class:"
    ws[f'A{row}'].font = Font(bold=True)
    ws[f'B{row}'] = student_data['student']['classes']['name']
    
    row += 1
    ws[f'A{row}'] = "Term:"
    ws[f'A{row}'].font = Font(bold=True)
    ws[f'B{row}'] = f"{student_data['term']['academic_year']} - {student_data['term']['term']}"
    
    row += 2
    
    # Scores table headers
    headers = ['Subject', 'CA1 (15)', 'CA2 (15)', 'Exam (70)', 'Total (100)', 'Grade', 'Remark']
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = border
    
    row += 1
    
    # Scores
    grand_total = 0
    for score in student_data['scores']:
        ws.cell(row=row, column=1, value=score['subjects']['name']).border = border
        ws.cell(row=row, column=2, value=float(score['ca1'])).border = border
        ws.cell(row=row, column=2).alignment = center_align
        ws.cell(row=row, column=3, value=float(score['ca2'])).border = border
        ws.cell(row=row, column=3).alignment = center_align
        ws.cell(row=row, column=4, value=float(score['exam'])).border = border
        ws.cell(row=row, column=4).alignment = center_align
        ws.cell(row=row, column=5, value=float(score['total'])).border = border
        ws.cell(row=row, column=5).alignment = center_align
        ws.cell(row=row, column=5).font = Font(bold=True)
        ws.cell(row=row, column=6, value=score['grade']).border = border
        ws.cell(row=row, column=6).alignment = center_align
        ws.cell(row=row, column=6).font = Font(bold=True)
        ws.cell(row=row, column=7, value=score['remark']).border = border
        
        grand_total += float(score['total'])
        row += 1
    
    # Summary
    row += 1
    ws.cell(row=row, column=1, value="GRAND TOTAL:").font = Font(bold=True)
    ws.cell(row=row, column=5, value=grand_total).font = Font(bold=True)
    
    row += 1
    ws.cell(row=row, column=1, value="AVERAGE:").font = Font(bold=True)
    ws.cell(row=row, column=5, value=student_data['statistics']['average']).font = Font(bold=True)
    
    row += 1
    ws.cell(row=row, column=1, value="POSITION:").font = Font(bold=True)
    position = student_data['statistics']['position']
    ws.cell(row=row, column=5, value=f"{position}{get_ordinal_suffix(position)}").font = Font(bold=True)
    
    # Auto-adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)
        
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        
        adjusted_width = min(max_length + 2, 40)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    # Save to BytesIO
    excel_file = io.BytesIO()
    wb.save(excel_file)
    excel_file.seek(0)
    
    return excel_file


def get_ordinal_suffix(n):
    """Get ordinal suffix for number (1st, 2nd, 3rd, etc.)"""
    if not n:
        return ''
    
    if 10 <= n % 100 <= 20:
        suffix = 'th'
    else:
        suffix = {1: 'st', 2: 'nd', 3: 'rd'}.get(n % 10, 'th')
    
    return suffix
